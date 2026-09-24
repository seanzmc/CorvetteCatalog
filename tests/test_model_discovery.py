"""Reproduction protects frozen observations and their deterministic serialization."""
import hashlib
import io
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tarfile
import tempfile
import unittest

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]


def normalized_bytes(content):
    # Check the scope before replacing timestamp values in the original bytes:
    # parsing and re-serializing here would conceal writer or key-order drift.
    def walk(value, parent=None):
        if isinstance(value, dict):
            for key, child in value.items():
                if key == 'submitted_at':
                    if parent != 'compact' or not isinstance(child, str):
                        raise AssertionError('Unexpected submitted_at outside compact')
                walk(child, key)
        elif isinstance(value, list):
            for child in value:
                walk(child, parent)
    walk(json.loads(content))
    return re.sub(rb'("submitted_at"\s*:\s*)"(?:[^"\\]|\\.)*"',
                  rb'\1"<timestamp>"', content)


class ModelDiscoveryReproductionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        (ROOT / '.local').mkdir(exist_ok=True)

    def test_price_anchors_identify_frozen_option_rows(self):
        with tarfile.open(ROOT / 'baselines/2026-09-06/workbook-runtime.tar.gz') as archive:
            workbook = load_workbook(io.BytesIO(archive.extractfile('stingray_master.xlsx').read()),
                                     data_only=True)
        for lane in ('zr1', 'zr1x'):
            accounting = json.loads((ROOT / f'docs/discovery/{lane}-accounting.json').read_text())
            sheet = workbook[f'{lane}_options']
            headers = [cell.value for cell in sheet[1]]
            rows = {row[headers.index('option_id')]: (number, dict(zip(headers, row)))
                    for number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2)
                    if any(value is not None for value in row)}
            self.assertEqual({price['option_id'] for price in accounting['option_prices']}, set(rows))
            for price in accounting['option_prices']:
                with self.subTest(lane=lane, option=price['option_id']):
                    number, row = rows[price['option_id']]
                    self.assertEqual(price['workbook_anchor'], f'{sheet.title}!A{number}:K{number}')
                    self.assertEqual(price['rpo'], row['rpo'])
                    self.assertEqual(price['baseline_amount'], row['price'])

    # The ZR1/ZR1X extractor reads the Git-ignored raw guide; CI declares it absent.
    @unittest.skipIf(os.environ.get('CATALOG_RAW_SOURCES') == 'absent',
                     'raw manufacturer guide is Git-ignored and absent (CATALOG_RAW_SOURCES=absent)')
    def test_extractors_reproduce_committed_bytes(self):
        for lane in ('zr1', 'zr1x'):
            with self.subTest(lane=lane), tempfile.TemporaryDirectory(dir=ROOT / '.local') as directory:
                result = subprocess.run(
                    [sys.executable, 'scripts/model_discovery.py', lane, directory],
                    cwd=ROOT, capture_output=True, text=True, timeout=120)
                self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
                for filename, folder in [(f'{lane}-structured-records.json', 'docs'),
                                         (f'{lane}-accounting.json', 'docs/discovery')]:
                    self.assertTrue((Path(directory) / filename).read_bytes()
                                    == (ROOT / folder / filename).read_bytes(),
                                    f'{filename}: extractor bytes differ')

    def test_every_lane_reproduces_committed_bytes(self):
        schema = json.loads((ROOT / 'docs/discovery/handoff-schema.json').read_text())
        with tempfile.TemporaryDirectory(dir=ROOT / '.local') as directory:
            # Only node and tar are available: an accidental git/reference-repo
            # dependency must fail even on the original author's workstation.
            tools = Path(directory) / 'bin'
            tools.mkdir()
            for name in ('node', 'tar'):
                executable = shutil.which(name)
                self.assertIsNotNone(executable, f'{name} is required')
                (tools / name).symlink_to(executable)
            # GNU tar (Linux CI) runs gzip for .tar.gz; bsdtar has it built in.
            if gzip := shutil.which('gzip'):
                (tools / 'gzip').symlink_to(gzip)
            environment = {**os.environ, 'PATH': str(tools), 'TMPDIR': directory}
            for lane in schema['lanes']['models']:
                with self.subTest(lane=lane):
                    result = subprocess.run(
                        ['node', 'scripts/model_discovery.mjs', lane, directory],
                        cwd=ROOT, env=environment, capture_output=True, text=True, timeout=180)
                    self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
                    filename = f'{lane}-runtime.json'
                    generated = normalized_bytes((Path(directory) / filename).read_bytes())
                    committed = normalized_bytes((ROOT / 'docs/discovery' / filename).read_bytes())
                    # Preserve historical provenance; verify both hashes before replacing
                    # only the probe identity in bytes (no JSON reserialization).
                    current_hash = hashlib.sha256((ROOT / 'scripts/model_discovery.mjs').read_bytes()).hexdigest()
                    historical_hash = 'a0e05ffe477b2eb42d2393f288b642b857c74e2b0261073cd61a8e3ceb44602d'
                    self.assertEqual(json.loads(generated)['provenance']['probe_sha256'], current_hash)
                    self.assertEqual(json.loads(committed)['provenance']['probe_sha256'], historical_hash)
                    generated = generated.replace(
                        f'"probe_sha256":"{current_hash}"'.encode(),
                        f'"probe_sha256":"{historical_hash}"'.encode(), 1)
                    self.assertTrue(generated == committed,
                                    f'{lane}: bytes differ beyond compact.submitted_at and verified probe identity')


if __name__ == '__main__':
    unittest.main()
