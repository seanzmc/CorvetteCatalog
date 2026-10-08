"""Read a whole manufacturer source and propose catalog updates that respect manual edits.

Run:
  python -m catalog.source_intake prices DRAFT NEW_SCHEDULE --previous OLD_SCHEDULE [--review OUT.json]
  python -m catalog.source_intake proposals DRAFT PROPOSALS.json [--review OUT.json]
  python -m catalog.source_intake apply DRAFT REVIEW.json [--keep ID ...] [--take ID ...]
  python -m catalog.source_intake accept DRAFT --reviewer NAME --reason TEXT

Each fact is compared three ways: the last manufacturer value, the new
manufacturer value and the current catalog value. A manufacturer change to a
fact nobody edited by hand is proposed as an update; a manual edit the
manufacturer did not change is kept; when both changed, the owner keeps the
manual value or takes the manufacturer's. Nothing is written until `apply`,
which saves ordinary reviewed edits and stages one intake assertion for each
model's price edit or each proposal, citing the exact source lines. `accept`
records the owner's acceptance; building a release remains a separate step.

Price schedules are read whole (the CSV a spreadsheet exports, or the order
guide workbook's Price Schedule sheet). Distribution updates are prose, so they
arrive as interpreted proposals: each cites its page and line, quotes the
source and lists the record edits it implies.
"""
import argparse
from contextlib import closing
import csv
import hashlib
import json
from pathlib import Path
import re
import shutil

from catalog import authoring, authoring_acceptance as acceptance, authoring_components, authoring_records as records
from catalog import authoring_relationships, foundation as f
from catalog.consumers import digest, encode
from catalog.releases import database_hash

INTAKE_REASON = 'Manufacturer intake: '
MODELS = {'C': 'stingray', 'E': 'grand_sport', 'G': 'grand_sport_x', 'H': 'z06', 'R': 'zr1', 'S': 'zr1x'}
MODEL_NAMES = [('Grand Sport X', 'grand_sport_x'), ('Grand Sport', 'grand_sport'), ('Stingray', 'stingray'),
               ('ZR1X', 'zr1x'), ('ZR1', 'zr1'), ('Z06', 'z06')]
# Words that may surround model names in a qualifier that only splits a price by model.
MODEL_ONLY = re.compile(r'(?:\s|,|&|/|\band\b|\bonly\b)*', re.I)


def money(value):
    """Minor units from '$1,295.00', '-$910.00' or a spreadsheet number; None otherwise."""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return round(value * 100)
    text = str(value or '').replace('$', '').replace(',', '').strip()
    return round(float(text) * 100) if re.fullmatch(r'-?\d+(?:\.\d+)?', text) else None


def preserve(path):
    """Copy a supplied original into Git-ignored sources/raw/<sha256>/, never overwriting."""
    path = Path(path).resolve()
    sha = hashlib.sha256(path.read_bytes()).hexdigest()
    target = f.ROOT / 'sources/raw' / sha / path.name
    if not target.exists():
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(path, target)
    if hashlib.sha256(target.read_bytes()).hexdigest() != sha:
        raise ValueError(f'A different file already occupies {target}')
    return target.relative_to(f.ROOT).as_posix(), sha


def _rows(path):
    """(locator, cells) for every nonblank row of a CSV export or the Price Schedule sheet."""
    path = Path(path)
    if path.suffix.lower() == '.xlsx':
        import openpyxl  # only needed for workbook exports
        sheet = openpyxl.load_workbook(path, data_only=True, read_only=True)['Price Schedule']
        for number, row in enumerate(sheet.iter_rows(values_only=True), 1):
            cells = ['' if v is None else v for v in row[:11]]
            if any(c != '' for c in cells):
                yield f'Price Schedule row {number}', cells + [''] * (11 - len(cells))
        return
    with path.open(newline='', encoding='utf-8-sig') as handle:
        reader, start = csv.reader(handle), 1
        for cells in reader:
            locator = f'line {start}' if reader.line_num == start else f'lines {start}-{reader.line_num}'
            start = reader.line_num + 1
            if any(c.strip() for c in cells):
                yield locator, [c.strip() for c in cells] + [''] * (11 - len(cells))


def read_price_schedule(path):
    """Every model and option price, with its exact source location.

    Both layouts carry List Price, D/H and MSRP in the same columns; the list
    price is used (owner decision ST-D12) and must equal MSRP minus D/H.
    """
    entries, section, category, effective, revised = [], None, '', None, None
    for locator, cells in _rows(path):
        first = str(cells[0])
        if first.upper().startswith('EFFECTIVE'):
            effective = first
        if re.search(r'Revised \w+ \d{1,2}, \d{4}', first):
            revised = re.search(r'Revised \w+ \d{1,2}, \d{4}', first).group(0)  # the year runs into the footer
        if first == 'Base Model Prices':
            section = 'model'; continue
        if first == 'Additional Options':
            section = 'option'; continue
        code = str(cells[1])
        if section == 'option' and code.endswith(':'):
            category = code.rstrip(':'); continue
        marked = first == '¨'
        if section == 'model' and re.fullmatch(r'1Y[A-Z]\d\d', code):
            amount, dh, msrp = money(cells[3]), money(cells[4]), money(cells[5])
            trim = str(cells[2]).split()[-1].lower()
            entries.append(dict(kind='model', code=code, description=str(cells[2]), qualifier='',
                                category='', amount_minor=amount, destination_minor=money(cells[9]),
                                model_key=MODELS[code[2]], configuration_id=f'{trim}_{code[2].lower()}{code[3:]}',
                                consistent=amount is not None and msrp == amount + (dh or 0),
                                marked=marked, locator=locator))
        elif section == 'option' and code and money(cells[4]) is not None:
            amount, dh, msrp = money(cells[4]), money(cells[5]), money(cells[6])
            entries.append(dict(kind='option', code=code, description=str(cells[2]), qualifier=str(cells[3]),
                                category=category, amount_minor=amount, consistent=msrp == amount + (dh or 0),
                                marked=marked, locator=locator))
    keys = [entry_key(e) for e in entries]
    if len(set(keys)) != len(keys):
        raise ValueError('Two price rows share a code and qualifier: ' + ', '.join(sorted({k for k in keys if keys.count(k) > 1})))
    if not any(e['kind'] == 'model' for e in entries):
        raise ValueError(f'{path} is not a price schedule (no base model prices)')
    return dict(effective=effective, revised=revised, entries=entries)


def entry_key(entry):
    if entry['kind'] == 'model':
        return 'model:' + entry['configuration_id']
    return f"option:{entry['code']}:{entry['qualifier']}"


def qualifier_models(qualifier):
    """Model keys a qualifier names, when it says nothing else; None for contextual prices."""
    found, rest = set(), qualifier
    for name, key in MODEL_NAMES:
        if re.search(r'\b' + re.escape(name) + r'\b', rest):
            found.add(key)
            rest = re.sub(r'\b' + re.escape(name) + r'\b', ' ', rest)
    return found if found and MODEL_ONLY.fullmatch(rest) else None


def revisions(db):
    return {r['model_key']: r['revision_id'] for r in authoring.catalog(db)}


def map_prices(db, schedule):
    """{entry key: [fact]}; contextual prices the catalog splits by condition are left unmapped."""
    revs, mapped = revisions(db), {}
    siblings = {}
    for entry in schedule['entries']:
        if entry['kind'] == 'option':
            siblings.setdefault(entry['code'], []).append(entry)
    for entry in schedule['entries']:
        key = entry_key(entry)
        if entry['kind'] == 'model':
            revision = revs[entry['model_key']]
            row = records.lookup(db, 'configuration', dict(revision_id=revision, id=entry['configuration_id']))
            if row is None:
                raise ValueError(f"No configuration {entry['configuration_id']} for {entry['description']}")
            mapped[key] = [dict(table='configuration', key=dict(revision_id=revision, id=row['id']),
                                field='starting_amount_minor', model_key=entry['model_key'], label=entry['description'],
                                value=entry['amount_minor'] + (entry['destination_minor'] or 0))]
            continue
        group = siblings[entry['code']]
        if any(e['amount_minor'] < 0 for e in group):
            mapped[key] = None  # discounts are folded into the catalog's contextual prices
            continue
        if len(group) == 1:
            models = set(revs)
        else:
            named = [qualifier_models(e['qualifier']) if e['qualifier'] else set() for e in group]
            if any(n is None for n in named) or sum(not n for n in named) > 1 or any(
                    a & b for i, a in enumerate(named) for b in named[i + 1:]):
                mapped[key] = None  # contextual price: the owner maps it to a contextual rate
                continue
            mine = named[group.index(entry)]
            models = mine or set(revs) - set().union(*named)
        facts = []
        for model in sorted(models):
            for row in db.execute('''SELECT * FROM option WHERE revision_id=? AND rpo=? AND charge_mode='priced'
                                     AND purchase_amount_minor IS NOT NULL ORDER BY id''', (revs[model], entry['code'])):
                contextual = db.execute('SELECT count(*) FROM option_rate WHERE revision_id=? AND target_option_id=?',
                                        (revs[model], row['id'])).fetchone()[0]
                facts.append(dict(table='option', key=dict(revision_id=revs[model], id=row['id']),
                                  field='purchase_amount_minor', model_key=model, label=f"{row['rpo']} {row['name']}",
                                  value=entry['amount_minor'], contextual_rates=contextual))
        mapped[key] = facts
    return mapped


def manual_edits(db):
    """{(table, key): [saved edit]} for edits a manufacturer intake did not apply."""
    present = acceptance.tables(db)
    linked = set()
    if 'authoring_intake' in present:
        for (ids,) in db.execute("SELECT accepted_change_ids_json FROM authoring_intake WHERE disposition='accepted'"):
            linked |= set(json.loads(ids or '[]'))
    result = {}
    for item in acceptance.events(db):
        event = item['event']
        if item['change_ref'] in linked or event['reason'].startswith(INTAKE_REASON):
            continue
        if item['table'] == 'authoring_record_change':
            keys = [(op['table'], op['key']) for op in json.loads(event['operations_json'])]
        elif item['table'] == 'authoring_change':
            keys = [('option', dict(revision_id=event['revision_id'], id=event['option_id']))]
        elif item['table'] == 'authoring_relationship_change':
            keys = [('acquisition', dict(revision_id=event['revision_id'], id=event['acquisition_id']))]
        else:
            keys = [('component_rate', dict(revision_id=event['revision_id'], component_id=event['component_id'],
                                            configuration_id=event['configuration_id']))]
        for table, key in keys:
            result.setdefault((table, encode(key)), []).append(
                dict(change_ref=item['change_ref'], saved_at=event['saved_at'], reason=event['reason']))
    return result


def _source(path):
    path = Path(path).resolve()
    try:
        relative = path.relative_to(f.ROOT).as_posix()
    except ValueError:
        relative = None
    return dict(name=path.name, path=relative, sha256=hashlib.sha256(path.read_bytes()).hexdigest())


def compare_prices(db, new_path, previous_path):
    """The review of a new price schedule against the previous one and the draft."""
    new, old = read_price_schedule(new_path), read_price_schedule(previous_path)
    new_map, old_map = map_prices(db, new), map_prices(db, old)
    old_entries = {entry_key(e): e for e in old['entries']}
    edits, items, notes = manual_edits(db), [], []
    for entry in new['entries']:
        key = entry_key(entry)
        before = old_entries.get(key)
        if not entry['consistent']:
            notes.append(dict(kind='inconsistent_row', locator=entry['locator'], code=entry['code'],
                              detail='List price plus D/H does not equal MSRP; the list price is used'))
        if before and before['description'] != entry['description']:
            notes.append(dict(kind='wording', locator=entry['locator'], code=entry['code'],
                              detail=f"{before['description']!r} is now {entry['description']!r}"))
        if before and before['amount_minor'] == entry['amount_minor'] and before.get('destination_minor') == entry.get('destination_minor'):
            continue  # no manufacturer change: any manual value stays
        if new_map[key] is None:
            notes.append(dict(kind='contextual_price', locator=entry['locator'], code=entry['code'],
                              detail=f"{entry['qualifier']}: {money_text(before and before['amount_minor'])} → "
                                     f"{money_text(entry['amount_minor'])}; map it to the contextual rate by hand"))
            continue
        if not new_map[key]:
            notes.append(dict(kind='not_in_catalog', locator=entry['locator'], code=entry['code'],
                              detail=f"{entry['description']} ({entry['qualifier'] or 'all models'}) "
                                     f"{money_text(entry['amount_minor'])} has no priced catalog option"))
            continue
        previous_facts = {encode(fact['key']): fact for fact in (old_map.get(key) or [])}
        for fact in new_map[key]:
            current = records.lookup(db, fact['table'], fact['key'])[fact['field']]
            last = previous_facts.get(encode(fact['key']), {}).get('value')
            manual = edits.get((fact['table'], encode(fact['key'])), [])
            if current == fact['value']:
                status = 'already_current'
            elif last is not None and current == last:
                status = 'update'
            elif last is None and not manual:
                status = 'update'
            else:
                status = 'conflict'
            item = dict(status=status, table=fact['table'], key=fact['key'], field=fact['field'],
                        model_key=fact['model_key'], label=fact['label'], last=last, new=fact['value'],
                        current=current, manual_edits=manual, locator=entry['locator'],
                        previous_locator=before and before['locator'], code=entry['code'],
                        contextual_rates=fact.get('contextual_rates', 0))
            item['id'] = digest(dict(source=_source(new_path)['sha256'], table=item['table'], key=item['key'],
                                     field=item['field']))[:12]
            items.append(item)
    for key, before in old_entries.items():
        if key not in {entry_key(e) for e in new['entries']}:
            notes.append(dict(kind='removed_row', locator=before['locator'], code=before['code'],
                              detail=f"{before['description']} ({before['qualifier'] or 'all models'}) is no longer listed; nothing is removed automatically"))
    return dict(format='source-intake-review-v1', kind='price_schedule', draft_etag=database_hash(db),
                source=_source(new_path) | dict(effective=new['effective'], revised=new['revised']),
                previous=_source(previous_path) | dict(effective=old['effective'], revised=old['revised']),
                items=items, notes=notes)


def compare_proposals(db, path):
    """The review of interpreted proposals (distribution updates) against the draft.

    Each proposal lists `expected` values it was interpreted against. If the
    draft now differs, or a manual edit touched one of its records, the owner
    keeps the catalog as it is or takes the proposal.
    """
    document = json.loads(Path(path).read_text())
    if document.get('format') != 'source-proposals-v1':
        raise ValueError('Unknown proposal format')
    revs, edits, items = revisions(db), manual_edits(db), []
    for proposal in document['proposals']:
        revision = revs[proposal['model_key']]
        differs = []
        for expected in proposal.get('expected', []):
            key = dict(revision_id=revision, **expected['key'])
            row = records.lookup(db, expected['table'], key)
            actual = None if row is None else {k: row[k] for k in expected['values']}
            if actual != expected['values']:
                differs.append(dict(table=expected['table'], key=key, expected=expected['values'], current=actual))
        manual = [e for step in proposal['steps'] for r in step
                  for e in edits.get((r['table'], encode(_target(revision, r))), [])]
        applied = all(_already(db, revision, request) for step in proposal['steps'] for request in step)
        status = 'already_current' if applied else 'conflict' if differs or manual else 'update'
        items.append(dict(id=proposal['id'], status=status, model_key=proposal['model_key'], label=proposal['summary'],
                          locator=proposal['locator'], quote=proposal['quote'], steps=proposal['steps'],
                          differs=differs, manual_edits=manual))
    return dict(format='source-intake-review-v1', kind='proposals', draft_etag=database_hash(db),
                source=dict(document['source'], proposals=_source(path)), items=items, notes=document.get('notes', []))


def _target(revision, request):
    """The record a request writes: a clone writes its new identity, not its source."""
    key = dict(revision_id=revision, **request['key'])
    return dict(key, id=request['new_id']) if request.get('action') == 'clone' else key


def _already(db, revision, request):
    action = request.get('action', 'update')
    row = records.lookup(db, request['table'], _target(revision, request))
    if action == 'delete':
        return row is None
    return row is not None and all(row[k] == v for k, v in request.get('values', {}).items())


def money_text(minor):
    if minor is None:
        return 'none'
    return ('-' if minor < 0 else '') + f'${abs(minor) // 100:,}' + (f'.{abs(minor) % 100:02d}' if minor % 100 else '')


def decide(review, keep=(), take=()):
    """Which items to apply: updates unless kept, conflicts only when taken."""
    known = {item['id']: item for item in review['items']}
    for identifier in (*keep, *take):
        if identifier not in known:
            raise ValueError(f'Unknown review item {identifier}')
    if set(keep) & set(take):
        raise ValueError('An item cannot be both kept and taken')
    undecided = [i for i, item in known.items() if item['status'] == 'conflict' and i not in keep and i not in take]
    if undecided:
        raise ValueError('Choose keep or take for every conflict: ' + ', '.join(undecided))
    return [item for i, item in known.items()
            if item['status'] != 'already_current' and i not in keep and (item['status'] == 'update' or i in take)]


def _batches(db, review, chosen):
    """(revision, [request lists], locators, label): price facts one edit per model, proposals one per step."""
    if review['kind'] == 'price_schedule':
        grouped = {}
        for item in chosen:
            requests, locators, labels = grouped.setdefault(item['key']['revision_id'], ([], [], []))
            requests.append(dict(table=item['table'], key=item['key'], values={item['field']: item['new']}))
            locators.append(item['locator'])
            labels.append(f"{item['model_key']} {item['label']} {money_text(item['current'])} → {money_text(item['new'])}")
        return [(rev, [requests], ', '.join(dict.fromkeys(locators)), '; '.join(labels))
                for rev, (requests, locators, labels) in grouped.items()]
    revs = revisions(db)
    return [(revs[item['model_key']],
             [[dict(r, key=dict(revision_id=revs[item['model_key']], **r['key'])) for r in step] for step in item['steps']],
             item['locator'], f"{item['model_key']}: {item['label']}. Source: {item['quote']}") for item in chosen]


def apply(db, review, keep=(), take=()):
    """Save each chosen batch as reviewed edits, then stage the intake assertion they cite.

    Steps already in the draft are skipped, so an interrupted apply can be rerun.
    """
    source = review['source']
    if not source.get('path') or not (review.get('previous') or source['proposals']).get('path'):
        raise ValueError('Preserve the sources in this repository (sources/raw) first')
    fresh = (compare_prices(db, f.ROOT / source['path'], f.ROOT / review['previous']['path'])
             if review['kind'] == 'price_schedule' else compare_proposals(db, f.ROOT / source['proposals']['path']))
    if fresh != review:
        raise ValueError('The draft or source changed since this review; review it again')
    chosen, saved = decide(review, keep, take), []
    for revision, steps, locator, label in _batches(db, review, chosen):
        payload = dict(
            source_sha256=source['sha256'], source_path=source['path'], locator=locator,
            assertion=label[:4000], scope=revision,
            comparison='added' if any(r.get('action') in ('clone', 'create') for s in steps for r in s) else 'changed',
            ambiguity=('none: list price column (ST-D12), destination added to base prices'
                       if review['kind'] == 'price_schedule' else 'interpreted from prose; confirmed by the owner at review'),
            coverage='complete' if review['kind'] == 'price_schedule' else 'partial')
        reason = f"{INTAKE_REASON}{digest(payload)[:12]} {source['name']} {locator}: {label}"[:2000]
        refs = []
        for requests in steps:
            if all(_already(db, revision, dict(r, key={k: v for k, v in r['key'].items() if k != 'revision_id'}))
                   for r in requests):
                continue
            try:
                result = records.save(db, records.preview(db, revision, database_hash(db), requests, reason))
            except ValueError as error:
                raise ValueError(f'{label[:120]} ({locator}) could not be saved: {error}. '
                                 f'Saved before it: {json.dumps(saved)}') from error
            refs.append(f"authoring_record_change:{revision}:{result['edit_version']}")
        staged = acceptance.stage(db, payload)
        saved.append(dict(intake_id=staged['intake_id'], change_refs=refs, locator=locator))
    return dict(applied=len(chosen), kept=sorted(keep), saved=saved, etag=database_hash(db))


def accept(db, reviewer, reason):
    """Accept the draft, linking each pending intake assertion to the edits that cite it."""
    history = acceptance.events(db)
    decisions = []
    for row in db.execute("SELECT intake_id FROM authoring_intake WHERE disposition IN ('pending','deferred') ORDER BY intake_id"):
        refs = [i['change_ref'] for i in history if i['event']['reason'].startswith(INTAKE_REASON + row['intake_id'][:12])]
        if refs:
            decisions.append(dict(intake_id=row['intake_id'], change_ids=refs,
                                  resolution='Applied through source intake after owner review'))
    change = acceptance.preview(db, database_hash(db), reviewer, reason, decisions)
    return acceptance.accept(db, change)


def open_draft(path):
    db = authoring.open_workspace(path)
    for module in (authoring_relationships, authoring_components, records, acceptance):
        module.prepare(db)
    return db


def render(review):
    """A short owner-facing summary of a review."""
    lines = [f"{review['source']['name']} ({review['source'].get('effective') or ''}"
             f"{', ' + review['source']['revised'] if review['source'].get('revised') else ''})"]
    if review['kind'] == 'price_schedule':
        lines.append(f"compared with {review['previous']['name']} ({review['previous'].get('revised') or ''})")
    for status in ('conflict', 'update', 'already_current'):
        chosen = [i for i in review['items'] if i['status'] == status]
        if not chosen:
            continue
        lines.append(f'\n{status.replace("_", " ").title()} ({len(chosen)}):')
        for item in chosen:
            if review['kind'] == 'price_schedule':
                lines.append(f"  [{item['id']}] {item['model_key']} {item['label']}: catalog {money_text(item['current'])}, "
                             f"manufacturer {money_text(item['last'])} → {money_text(item['new'])} ({item['locator']})")
            else:
                lines.append(f"  [{item['id']}] {item['model_key']} {item['label']} ({item['locator']})")
            for edit in item['manual_edits']:
                lines.append(f"      manual edit {edit['saved_at'][:10]}: {edit['reason'][:120]}")
    if review['notes']:
        lines.append(f"\nNotes ({len(review['notes'])}):")
        lines += [f"  {n['kind']} {n['code'] if 'code' in n else ''} ({n.get('locator', '')}): {n['detail']}" for n in review['notes']]
    return '\n'.join(lines)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    commands = parser.add_subparsers(dest='command', required=True)
    prices = commands.add_parser('prices', help='Review a new price schedule')
    prices.add_argument('draft', type=Path)
    prices.add_argument('schedule', type=Path)
    prices.add_argument('--previous', type=Path, required=True, help='The last price schedule taken into the catalog')
    proposals = commands.add_parser('proposals', help='Review interpreted proposals (distribution updates)')
    proposals.add_argument('draft', type=Path)
    proposals.add_argument('proposals', type=Path)
    for command in (prices, proposals):
        command.add_argument('--review', type=Path, help='Where to write the review JSON')
    run = commands.add_parser('apply', help='Save the reviewed changes as edits awaiting acceptance')
    run.add_argument('draft', type=Path)
    run.add_argument('review', type=Path)
    run.add_argument('--keep', nargs='*', default=[], help='Items to leave as they are in the catalog')
    run.add_argument('--take', nargs='*', default=[], help='Conflicts to overwrite with the manufacturer value')
    done = commands.add_parser('accept', help='Accept the draft and link intake assertions to their edits')
    done.add_argument('draft', type=Path)
    done.add_argument('--reviewer', required=True)
    done.add_argument('--reason', required=True)
    args = parser.parse_args()
    with closing(open_draft(args.draft)) as db:
        if args.command == 'apply':
            print(json.dumps(apply(db, json.loads(args.review.read_text()), args.keep, args.take), indent=2))
            return
        if args.command == 'accept':
            print(json.dumps(accept(db, args.reviewer, args.reason), indent=2))
            return
        if args.command == 'prices':
            review = compare_prices(db, f.ROOT / preserve(args.schedule)[0], f.ROOT / preserve(args.previous)[0])
        else:
            review = compare_proposals(db, args.proposals)
    if args.review:
        args.review.write_text(json.dumps(review, indent=2, sort_keys=True) + '\n')
    print(render(review))


if __name__ == '__main__':
    main()
