import io, json, tarfile, hashlib
from pathlib import Path
from openpyxl import load_workbook

# Read-only source extraction. Usage: python scripts/model_discovery.py LANE NEW_OUTPUT_DIR
# Raw input bytes stay immutable; every retained row keeps its original Excel locator.

import sys
lane=sys.argv[1]
assert lane in {'zr1'}, 'Supply a supported lane'
root=Path(__file__).resolve().parents[1]; out=Path(sys.argv[2]).resolve()
assert not (out/f'{lane}-structured-records.json').exists(), 'Choose a fresh output directory'
out.mkdir(parents=True,exist_ok=True)
manifest=json.loads((root/'baselines/2026-09-06/manifest.json').read_text())
with tarfile.open(root/'baselines/2026-09-06/workbook-runtime.tar.gz') as t:
    for name in ['stingray_master.xlsx','form-app/app.js','form-app/data.js',f'form-output/runtime/{lane}-runtime-contract.json']:
        b=t.extractfile(name).read()
        assert hashlib.sha256(b).hexdigest()==next(x['sha256'] for x in manifest['files'] if x['path']==name)
        if name.endswith('.xlsx'): w=load_workbook(io.BytesIO(b),data_only=True)
        else:
            dest=out/name; dest.parent.mkdir(parents=True,exist_ok=True); dest.write_bytes(b)
guide_sha256='d3ca7d3a09c9fb89210b4ce584493b3ad8fb65ca35087c49d816d1cbf1a333d1'
raw=next((root/'sources/raw'/guide_sha256).glob('*.xlsx'))
assert hashlib.sha256(raw.read_bytes()).hexdigest()==guide_sha256
g=load_workbook(raw,rich_text=True,data_only=True)
records={}
for s in w:
    headers=[c.value for c in s[1]]
    records[s.title]=[dict(_row=i,**{str(k):v for k,v in zip(headers,row) if k is not None}) for i,row in enumerate(s.iter_rows(min_row=2,values_only=True),2) if any(v is not None for v in row)]


from collections import Counter,defaultdict
import re
options=records[f'{lane}_options']; bycode=defaultdict(list)
for r in options:
 if r['rpo']:bycode[r['rpo']].append(r)
ovs={(r['option_id'],r['variant_id']):r for r in records[f'{lane}_ovs']}
variants=['1lz_r07','3lz_r07','1lz_r67','3lz_r67']
def symbol(cell):
 v=cell.value
 if isinstance(v,str):return re.sub(r'\d+','',v).strip()
 return ''.join(str(x) for x in v if isinstance(x,str) or not x.font.vertAlign).strip()
primary=[]; missing=[]; mismatches=[]; pairs=0; coded={}
for s in ['Interior 5','Exterior 5','Mechanical 5']:
 for row in g[s].iter_rows(min_row=4):
  code=str(row[0].value or row[1].value or '')
  if not code:continue
  if all(symbol(c)=='--' for c in row[3:7]):continue
  syms=[symbol(c) for c in row[3:7]]
  coded[code]=(s,row[0].row,syms,str(row[2].value))
  if code not in bycode:missing.append([s,row[0].row,code,str(row[2].value)]);continue
  opt=next((o for o in bycode[code] if o['active']),bycode[code][0]); differences=[]
  for v,marker in zip(variants,syms):
   status={'S':'standard','■':'standard','□':'standard','A':'available','A/D':'available','--':'unavailable'}[marker]
   wr=ovs[(opt['option_id'],v)];pairs+=1
   if wr['status']!=status:differences.append([v,status,wr['status'],wr['_row']])
  entry=dict(guide=f'{s}!A{row[0].row}:G{row[0].row}',code=code,workbook_row=opt['_row'],active=opt['active'],symbols=syms,differences=differences)
  primary.append(entry)
  if differences:mismatches.append(entry)
extras=[{'row':r['_row'],'code':r['rpo'],'name':r['option_name']} for r in options if r['rpo'] not in coded]
duplicates=[]
for s in ['Standard Equipment 5','Equipment Groups 5']:
 for row in g[s].iter_rows(min_row=4):
  code=str(row[0].value or row[1].value or '')
  if code not in coded or all(symbol(c)=='--' for c in row[3:7]):continue
  source,n,syms,desc=coded[code]
  duplicates.append(dict(guide=f'{s}!A{row[0].row}:G{row[0].row}',code=code,same_text=str(row[2].value)==desc,same_status=[symbol(c) for c in row[3:7]]==syms))
report=dict(primary=primary,pairs=pairs,missing=missing,status_differences=mismatches,workbook_only_or_uncoded=extras,duplicate_occurrences=duplicates)
(out/'remaining-coverage.json').write_text(json.dumps(report,indent=2))

from collections import Counter,defaultdict
def plain(cell):
 v=cell.value
 if isinstance(v,str) or v is None:return v
 return ''.join(str(x) for x in v if isinstance(x,str) or not x.font.vertAlign)
ids={r['interior_id'] for r in records['model_interior_scope'] if r['model_key']==lane}
ints=[r for r in records['LZ_Interiors'] if r['interior_id'] in ids]
components=defaultdict(list)
for r in records['interior_components']:
 if r['model_key']==lane:components[r['interior_id']].append(r)
base={};raw_combos=set()
for sheet,rows,cols in [('Color and Trim 1',range(5,13),range(5,18)),('Color and Trim 2',range(5,9),range(5,9))]:
 for n in rows:
  s=g[sheet];trim=str(s.cell(n,1).value).split(',')[-1].strip()
  if trim=='2LZ':continue
  seats=str(s.cell(n,3).value).split(' / ')
  for col in cols:
   code=plain(s.cell(n,col))
   if not code or code=='--':continue
   for seat in seats:base[(trim,seat,code)]=(sheet,n,col);raw_combos.add((trim,seat,code))
actual_combos={(r['Trim'].replace('_R6X',''),r['Seat'],r['Interior Code']) for r in ints}
issues=[];interior_records=[];color_expected=set()
stitch={'36S':{'H1Y','HTM','HTP','HTE','HTT','HUB','HUC'},'37S':{'H1Y','HTM','HTP','HTE','HTT','HUB','HUC'},'38S':{'H1Y','HTM','HTP','HTE','HTT','HUB','HUC','HU0','HXO'}}
two={'HUK','HU6','HUL','HU7','HTN','HTQ','HU1','HU9','HU2','HUA','HUE','HTG','HMO','HVV','HU0','HXO'}
for r in ints:
 key=(r['Trim'].replace('_R6X',''),r['Seat'],r['Interior Code']); sheet,n,col=base[key];code=r['Interior Code']; cs=components[r['interior_id']]; codes={c['rpo'] for c in cs}
 problems=[]
 if r['Suede'] in {'N26','N2Z'} and r['Suede'] not in codes:problems.append('missing N26 component')
 if r['Stitch'] and (r['Stitch'] not in codes or code not in stitch[r['Stitch']]):problems.append('stitch mismatch')
 if r['Two Tone'] and (r['Seat']!='AH2' or code not in two or 'TU7' not in codes):problems.append('two tone mismatch')
 if r['Seat']=='AH2' and code in {'HU7','HUA','HXO'} and 'TU7' not in codes:problems.append('missing mandatory TU7')
 if code in {'HU0','HXO'} and '38S' not in codes:problems.append('missing mandatory 38S')
 if (sheet=='Color and Trim 2') != ('R6X' in codes):problems.append('R6X scope mismatch')
 if problems:issues.append([r['interior_id'],problems])
 s=g[sheet]
 for pn in (range(16,26) if sheet.endswith('1') else range(11,21)):
  if plain(s.cell(pn,col))=='--':color_expected.add((r['interior_id'],str(s.cell(pn,3).value)))
 interior_records.append({'id':r['interior_id'],'workbook':f"LZ_Interiors!A{r['_row']}:Q{r['_row']}",'guide':f'{sheet}!{s.cell(n,col).coordinate}','components':[c['rpo'] for c in cs]})
paint_codes={r['option_id']:r['rpo'] for r in records[f'{lane}_options'] if r['rpo'] in ['G26','G4Z','G8G','GBA','GBK','GEC','GKA','GKZ','GPH','GTR']}
actual={(r['interior_id'],paint_codes[r['option_id']]) for r in records['color_overrides'] if r['interior_id'] in ids and r['option_id'] in paint_codes}
expected_leaves=set()
for (trim,seat,code),(sheet,n,col) in base.items():
 required=[]
 if n in ([6,9,11] if sheet.endswith('1') else [7,8]):required.append('N2Z' if trim=='3LZ' else 'N26')
 if sheet.endswith('2'):required.append('R6X')
 stitch_choices=[None]
 if code in {'HU0','HXO'}:stitch_choices=['38S']
 elif code in stitch['36S']:stitch_choices=[None,'36S','37S','38S']
 tones=[None]
 if seat=='AH2' and code in two:tones=['TU7'] if code in {'HU7','HUA','HXO'} else [None,'TU7']
 for sc in stitch_choices:
  for tc in tones:expected_leaves.add((trim,seat,code,tuple(sorted(required+[x for x in [sc,tc] if x]))))
actual_leaves={(r['Trim'].replace('_R6X',''),r['Seat'],r['Interior Code'],tuple(sorted(c['rpo'] for c in components[r['interior_id']] if c['component_type']!='seat'))) for r in ints}
(out/'interior-reconciliation.json').write_text(json.dumps(dict(interiors=interior_records,component_issues=issues,color_expected=sorted(color_expected),color_missing=sorted(color_expected-actual),color_extra=sorted(actual-color_expected)),indent=2))

assert not mismatches and not issues
assert raw_combos==actual_combos and expected_leaves==actual_leaves
assert color_expected==actual
assert all(x['same_text'] and x['same_status'] for x in duplicates)
variant_ids={r['variant_id'] for r in records['model_variants'] if r['model_key']==lane}
option_ids={r['option_id'] for r in options}
baseline={k:v for k,v in records.items() if k.startswith(f'{lane}_')}
for name in ['model_master','model_registry_promotion','model_workbook_sources','model_variants','context_section_master','section_presentation','runtime_steps','order_summary_sections','step_order_summary_map','default_selection_rules','model_interior_scope','interior_components']:
 baseline[name]=[r for r in records[name] if r['model_key']==lane]
baseline['variant_master']=[r for r in records['variant_master'] if r['variant_id'] in variant_ids]
baseline['LZ_Interiors']=ints
section_ids={r['section_id'] for r in options+ints+baseline['section_presentation']}
baseline['section_master']=[r for r in records['section_master'] if r['section_id'] in section_ids]
for name in ['context_choice_copy','asset_map']:baseline[name]=[r for r in records[name] if r['model_key'] in (lane,'*')]
for name in ['PriceRef','rule_phrase_map','runtime_rule_exceptions']:baseline[name]=records[name]
baseline['color_overrides']=[r for r in records['color_overrides'] if r['interior_id'] in ids and r['option_id'] in option_ids]
# Explicit, inspected uncoded equipment correspondence; no fuzzy-name acceptance.
uncoded={25:('Interior 5',89),147:('Interior 5',6),148:('Interior 5',16),149:('Interior 5',17),150:('Interior 5',19),151:('Interior 5',20),152:('Interior 5',22),153:('Interior 5',24),154:('Interior 5',26),155:('Interior 5',59),156:('Interior 5',84),157:('Interior 5',94),158:('Interior 5',98),159:('Interior 5',101),160:('Interior 5',102),161:('Exterior 5',4),162:('Exterior 5',12),163:('Exterior 5',16),164:('Exterior 5',75),165:('Mechanical 5',4),166:('Mechanical 5',10),167:('Mechanical 5',34),168:('Mechanical 5',39),169:('Mechanical 5',42),170:('Mechanical 5',47),197:('Interior 5',51)}
guide_rows=[]
for sn in ['Interior 5','Exterior 5','Mechanical 5','Standard Equipment 5','Equipment Groups 5']:
 for rr in g[sn].iter_rows(min_row=4):
  if rr[2].value is None:continue
  guide_rows.append(dict(anchor=f'{sn}!A{rr[0].row}:K{rr[0].row}',sheet=sn,row=rr[0].row,orderable_code=str(rr[0].value or ''),reference_code=str(rr[1].value or ''),disclosure=str(rr[2].value),zr1_symbols=[symbol(c) for c in rr[3:7]],zr1x_symbols=[symbol(c) for c in rr[7:11]],applicable_to_zr1=any(symbol(c)!='--' for c in rr[3:7])))
dispositions=[]
for o in options:
 matches=[x for x in primary if x['workbook_row']==o['_row']]
 anchors=[x['guide'] for x in matches]
 classification='coded_guide_match'
 if o['_row'] in uncoded:
  sn,n=uncoded[o['_row']];anchors=[f'{sn}!A{n}:G{n}'];classification='uncoded_equipment_match'
  for v,c in zip(variants,list(g[sn][n])[3:7]):
   assert ovs[(o['option_id'],v)]['status']=={'S':'standard','■':'standard','□':'standard','A':'available','A/D':'available','--':'unavailable'}[symbol(c)]
 elif not anchors:
  assert o['rpo'] in set(paint_codes.values()),o
  classification='color_chart';anchors=['Color and Trim 1!A16:Q25','Color and Trim 2!A11:H20']
 disclosures=[r['disclosure'] for r in guide_rows if any(a.split('!')[0]==r['sheet'] and int(re.search(r'!A(\d+)',a).group(1))==r['row'] for a in anchors if re.search(r'!A(\d+)',a))]
 assert len(disclosures)<=1,o
 dispositions.append(dict(record_id=f'{lane}:'+o['option_id'],workbook_row=o['_row'],rpo=o['rpo'],source_classification=classification,guide_anchors=anchors,guide_disclosure=disclosures[0] if disclosures else None))
price_rows={}
for rr in g['Price Schedule'].iter_rows(min_row=48):
 if rr[1].value:
  price_rows.setdefault(str(rr[1].value).strip(),[]).append(dict(anchor=f'Price Schedule!B{rr[0].row}:E{rr[0].row}',row=rr[0].row,description=rr[2].value,qualifier=rr[3].value,amount=rr[4].value))
STANDARD_CODES=['B6P','ZZ3','D3V','SL9','DY0','CFV','C2Z','CFC','AH2']
EXPLANATION={'null_baseline_not_an_inferred_zero':'Preserve null source amount; scope/standard equipment and accepted decisions decide any future charge. Null is not zero.',
 'source_rate_match':'Matching amount is source evidence; it is not proof of resolved build totals.',
 'source_rate_match_with_context':'Matching amount is source evidence; candidate qualifiers and conditional-price rows retain context. It is not proof of resolved build totals.',
 'zero_without_schedule_rate':'Known baseline zero for standard/default/paint/inactive choices without a schedule rate; not a discovered price.',
 'standard_equipment_not_purchase':'ZR1 model/body/trim standard equipment; a same-code paid rate for another model is not a ZR1 purchase.'}
prices=[]
for o in options:
 value=o['price'];candidates=price_rows.get(o['rpo'],[]) if o['rpo'] else []
 if value is None:classification='null_baseline_not_an_inferred_zero'
 elif o['rpo'] in STANDARD_CODES:classification='standard_equipment_not_purchase'
 elif any(r['amount']==value for r in candidates):classification='source_rate_match_with_context' if any(r['qualifier'] for r in candidates) else 'source_rate_match'
 elif value==0 and not candidates:classification='zero_without_schedule_rate'
 else:classification='unresolved_rate_difference'
 prices.append(dict(option_id=o['option_id'],rpo=o['rpo'],workbook_anchor=f"zr1_options!A{o['_row']}:K{o['_row']}",baseline_amount=value,classification=classification,explanation=EXPLANATION.get(classification),source_rates=candidates,conditional_price_rows=[r['_row'] for r in baseline[f'{lane}_price_rules'] if r['target_option_id']==o['option_id']]))
assert not [r for r in prices if r['classification']=='unresolved_rate_difference']
contract=json.loads((out/f'form-output/runtime/{lane}-runtime-contract.json').read_text())
triplet=lambda r:(r['source_id'],r['rule_type'],r['target_id'])
translations=[]
for r in baseline[f'{lane}_rule_mapping']:
 emitted=[x for x in contract['rules'] if triplet(x)==triplet(r)]
 inactive=[o['option_id'] for o in options if not o['active'] and o['option_id'] in (r['source_id'],r['target_id'])]
 assert emitted or inactive,r
 translations.append(dict(source_row=r['_row'],rule_id=r['rule_id'],source_id=r['source_id'],rule_type=r['rule_type'],target_id=r['target_id'],runtime_disposition='emitted' if emitted else 'inactive_endpoint_filtered',inactive_endpoints=inactive,runtime_rows=emitted))
raw_edges={triplet(r) for r in baseline[f'{lane}_rule_mapping']}
derived=[r for r in contract['rules'] if triplet(r) not in raw_edges]
missing_records=[dict(guide_anchor=f'{sn}!A{n}:G{n}',rpo=c,guide_disclosure=t,source_classification='omitted_offering' if c=='SAI' else 'interior_component' if c in ['TU7','N26','N2Z','36S','37S','38S'] else 'outside_customer_selection_scope') for sn,n,c,t in missing]
sheet_roles=dict(options=f'{lane}_options',availability=f'{lane}_ovs',variant_overrides=f'{lane}_variant_overrides',rule_mapping=f'{lane}_rule_mapping',rule_groups=f'{lane}_rule_groups',rule_group_members=f'{lane}_rule_group_members',exclusive_groups=f'{lane}_exclusive_groups',exclusive_group_members=f'{lane}_exclusive_members',price_rules=f'{lane}_price_rules',interiors='LZ_Interiors',color_overrides='color_overrides')
assert all(v in baseline for v in sheet_roles.values())
# Contract shapes: docs/discovery/handoff-schema.json. Owner decisions are never generated here.
source=dict(format='model-review-records-v2',model_key=lane,model_year=2027,role='Frozen source and baseline evidence for model discovery; accepted targets live in the owner-decisions overlay, not here',sources=dict(workbook_sha256=manifest['files'][next(i for i,r in enumerate(manifest['files']) if r['path']=='stingray_master.xlsx')]['sha256'],guide_sha256=hashlib.sha256(raw.read_bytes()).hexdigest(),runtime_commit=manifest['reference_commit'],workbook_row_locator='baseline_rows key is source sheet; _row is original Excel row; fields retain source headers and nulls',guide_model_columns='D:G only; H:K retained as ZR1X context, not ZR1 facts',observed_provenance=None),sheet_roles=sheet_roles,baseline_rows=baseline,offering_dispositions=dispositions,guide_only_dispositions=missing_records,interior_source_links=interior_records,duplicate_rpos_within_model={c:[o['option_id'] for o in options if o['rpo']==c] for c,n in Counter(o['rpo'] for o in options if o['rpo']).items() if n>1},runtime_derived_relationships=dict(role='Frozen emitted relationships absent from direct workbook rows; explicit future translation ownership required, not new owner corrections',workbook_direct_count=len(translations),emitted_direct_count=sum(1 for r in translations if r['runtime_disposition']=='emitted')+len(derived),records=derived),source_reconciliation=dict(primary_status_comparison=primary,repeated_guide_occurrences=duplicates,interior_reconciliation=json.loads((out/'interior-reconciliation.json').read_text()),base_prices=[dict(variant_id=v['variant_id'],workbook_row=v['_row'],base=v['base_price'],guide_anchor=f'Price Schedule!F{n}:J{n}',guide_base=g['Price Schedule'].cell(n,6).value+g['Price Schedule'].cell(n,10).value) for v,n in zip(baseline['variant_master'],[34,36,35,37])]),guide_rows=guide_rows)
accounting=dict(format='model-discovery-accounting-v1',model_key=lane,role='Supplemental discovery accounting, not accepted replacement data',guide_sha256=source['sources']['guide_sha256'],option_prices=prices,direct_rule_translation=translations,existing_decision_document=f'../{lane}-structured.md#8-decision-overlay-source-baseline-and-target-remain-separate')
assert all(x['base']==x['guide_base'] for x in source['source_reconciliation']['base_prices'])
# Preserve the committed contract migration's top-level field order.
source['runtime_derived_relationships']=source.pop('runtime_derived_relationships')
# One record per line keeps evidence reviewable without changing JSON semantics.
def encode(v,level=0):
 pad='  '*level
 if isinstance(v,dict) and any(isinstance(x,(dict,list)) for x in v.values()):return '{\n'+',\n'.join(pad+'  '+json.dumps(k)+': '+encode(x,level+1) for k,x in v.items())+'\n'+pad+'}'
 if isinstance(v,list):return '[\n'+',\n'.join(pad+'  '+json.dumps(x,ensure_ascii=False) for x in v)+'\n'+pad+']' if v else '[]'
 return json.dumps(v,ensure_ascii=False)
(out/f'{lane}-structured-records.json').write_text(encode(source)+'\n')
(out/f'{lane}-accounting.json').write_text(encode(accounting)+'\n')
print('Verified source accounting:',len(options),'offerings;',len(primary)*4,'coded status pairs;',len(duplicates),'repeated rows;',len(ints),'interiors;',len(translations),'direct rules;',len(derived),'derived rules')
